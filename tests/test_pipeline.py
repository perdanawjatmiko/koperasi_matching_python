from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from koperasi_match.config import MatchConfig, timestamped_output_path
from koperasi_match.pipeline import run
from koperasi_match.workbook_reader import province_from_a1


def _save_source(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "ACEH"
    ws.append(["PROVINSI ACEH"])
    ws.append(["NO", "ALAMAT"])
    ws.append([1, "LAYEUN, LEUPUNG, ACEH BESAR"])
    wb.save(path)


def _save_multi_sheet_source(path: Path):
    wb = Workbook()
    first = wb.active
    first.title = "SUMBER_1"
    first.append(["PROVINSI ACEH"])
    first.append(["NO", "ALAMAT"])
    first.append([1, "LAYEUN, LEUPUNG, ACEH BESAR"])
    second = wb.create_sheet("SUMBER_2")
    second.append(["PROVINSI ACEH"])
    second.append(["NO", "ALAMAT"])
    second.append([2, "LAYEUN, LEUPUNG, ACEH BESAR"])
    ignored = wb.create_sheet("CATATAN")
    ignored.append(["Tidak memiliki kolom alamat"])
    wb.save(path)


def _save_reference(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Export Laporan"
    ws.append(["NIK", "desa", "kecamatan", "Kota/Kabupaten", "Provinsi", "pic phone"])
    ws.append(["001", "Layeun", "Leupung", "Kabupaten Aceh Besar", "Aceh (NAD)", "0812"])
    wb.save(path)


def test_pipeline_preserves_inputs_and_identifiers(tmp_path):
    source = tmp_path / "source.xlsx"
    reference = tmp_path / "reference.xlsx"
    output = tmp_path / "output.xlsx"
    _save_source(source)
    _save_reference(reference)
    source_before = source.read_bytes()
    reference_before = reference.read_bytes()
    counts = run(MatchConfig(source=source, reference=reference, output=output, sheets=["ACEH"]))
    assert counts["MATCHED_EXACT"] == 1
    assert source.read_bytes() == source_before
    assert reference.read_bytes() == reference_before
    wb = load_workbook(output, data_only=True)
    assert wb["MATCHED"]["A2"].value == "001"
    assert wb["MATCHED"]["F2"].value == "0812"
    assert wb.sheetnames == ["RINGKASAN", "MATCHED", "PERLU_REVIEW", "TIDAK_DITEMUKAN"]


def test_province_is_read_only_from_a1():
    wb = Workbook()
    ws = wb.active
    ws.title = "NAMA_SHEET_BEBAS"
    ws["A1"] = "PROVINSI SUMATRA UTARA"
    ws["B2"] = "ACEH"
    assert province_from_a1(ws) == "SUMATERA UTARA"


def test_empty_a1_is_rejected():
    wb = Workbook()
    ws = wb.active
    ws.title = "ACEH"
    try:
        province_from_a1(ws)
    except ValueError as error:
        assert "sel A1" in str(error)
    else:
        raise AssertionError("A1 kosong harus menghasilkan ValueError")


def test_output_name_gets_timestamp_and_sequence(tmp_path):
    base = tmp_path / "hasil_pencocokan.xlsx"
    moment = datetime(2026, 9, 10, 14, 30, 25)
    first = timestamped_output_path(base, moment)
    assert first.name == "hasil_pencocokan_20260910_143025.xlsx"
    first.touch()
    second = timestamped_output_path(base, moment)
    assert second.name == "hasil_pencocokan_20260910_143025_2.xlsx"


def test_empty_sheet_option_processes_all_address_sheets_and_reports_progress(tmp_path):
    source = tmp_path / "source.xlsx"
    reference = tmp_path / "reference.xlsx"
    output = tmp_path / "output.xlsx"
    _save_multi_sheet_source(source)
    _save_reference(reference)
    events = []
    counts = run(
        MatchConfig(source=source, reference=reference, output=output, sheets=[]),
        progress=lambda message, current, total: events.append((message, current, total)),
    )
    assert counts["MATCHED_EXACT"] == 2
    assert any(
        message == "Mencocokkan data" and current == 2 and total == 2
        for message, current, total in events
    )
    workbook = load_workbook(output, data_only=True)
    assert workbook["MATCHED"].max_row == 3
    assert workbook["RINGKASAN"]["B5"].value == "SUMBER_1, SUMBER_2"
