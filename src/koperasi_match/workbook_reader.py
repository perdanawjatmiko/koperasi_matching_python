from collections.abc import Iterable
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from .models import ReferenceRecord, SourceRecord
from .normalizer import clean_spaces, normalize_header, normalize_kabupaten, normalize_name

REQUIRED_REFERENCE_HEADERS = {
    "desa": "DESA",
    "kecamatan": "KECAMATAN",
    "kabupaten": "KOTAKABUPATEN",
    "provinsi": "PROVINSI",
}


def _header_map(values: Iterable[object]) -> dict[str, int]:
    return {
        normalize_header(value): index for index, value in enumerate(values) if value is not None
    }


def find_header(sheet: Worksheet, target: str, scan_rows: int = 20) -> tuple[int, int]:
    wanted = normalize_header(target)
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, scan_rows)):
        for cell in row:
            if normalize_header(cell.value) == wanted:
                return cell.row, cell.column
    raise ValueError(f"Kolom {target!r} tidak ditemukan pada sheet {sheet.title!r}")


def province_from_a1(sheet: Worksheet) -> str:
    """Ambil dan normalisasi nama provinsi hanya dari sel A1."""
    raw_province = clean_spaces(sheet["A1"].value)
    if not raw_province:
        raise ValueError(f"Nama provinsi pada sel A1 sheet {sheet.title!r} kosong")
    province = normalize_name(raw_province, kind="province")
    if not province:
        raise ValueError(
            f"Nama provinsi pada sel A1 sheet {sheet.title!r} tidak valid: {raw_province!r}"
        )
    return province


def read_sources(path: Path, sheets: list[str], address_column: str) -> list[SourceRecord]:
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    missing = [name for name in sheets if name not in workbook.sheetnames]
    if missing:
        raise ValueError(f"Sheet sumber tidak ditemukan: {', '.join(missing)}")
    records: list[SourceRecord] = []
    for sheet_name in sheets:
        sheet = workbook[sheet_name]
        header_row, address_col = find_header(sheet, address_column)
        province = province_from_a1(sheet)
        for row_number in range(header_row + 1, sheet.max_row + 1):
            value = sheet.cell(row_number, address_col).value
            address = clean_spaces(value)
            if not address or address.upper().startswith("JUMLAH"):
                continue
            records.append(SourceRecord(sheet_name, row_number, address, province))
    workbook.close()
    return records


def discover_source_sheets(path: Path, address_column: str) -> list[str]:
    """Temukan seluruh sheet yang memiliki header kolom alamat."""
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    selected = []
    for sheet in workbook.worksheets:
        try:
            find_header(sheet, address_column)
        except ValueError:
            continue
        selected.append(sheet.title)
    workbook.close()
    if not selected:
        raise ValueError(f"Tidak ada sheet dengan kolom {address_column!r}")
    return selected


def read_reference(
    path: Path, sheet_name: str
) -> tuple[list[object], list[ReferenceRecord], dict[str, int]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet referensi tidak ditemukan: {sheet_name}")
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    mapping = _header_map(headers)
    missing = [
        label
        for label, normalized in REQUIRED_REFERENCE_HEADERS.items()
        if normalized not in mapping
    ]
    if missing:
        raise ValueError(f"Kolom referensi wajib tidak ditemukan: {', '.join(missing)}")
    indexes = {label: mapping[key] for label, key in REQUIRED_REFERENCE_HEADERS.items()}
    nik_index = mapping.get("NIK", -1)
    phone_index = mapping.get("PICPHONE", -1)
    records = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        mutable_values = list(row[: len(headers)])
        for identifier_index in (nik_index, phone_index):
            if identifier_index >= 0 and mutable_values[identifier_index] is not None:
                mutable_values[identifier_index] = str(mutable_values[identifier_index])
        values = tuple(mutable_values)
        if not any(value not in (None, "") for value in values):
            continue
        desa_value = values[indexes["desa"]]
        kec_value = values[indexes["kecamatan"]]
        kab_value = values[indexes["kabupaten"]]
        province_value = values[indexes["provinsi"]]
        records.append(
            ReferenceRecord(
                row=row_number,
                values=values,
                nik=str(values[nik_index] or "") if nik_index >= 0 else "",
                desa_original=clean_spaces(desa_value),
                kecamatan_original=clean_spaces(kec_value),
                kabupaten_original=clean_spaces(kab_value),
                province_original=clean_spaces(province_value),
                desa=normalize_name(desa_value, kind="desa"),
                kecamatan=normalize_name(kec_value, kind="kecamatan"),
                kabupaten=normalize_kabupaten(kab_value),
                province=normalize_name(province_value, kind="province"),
            )
        )
    workbook.close()
    return headers, records, indexes


def inspect_workbook(path: Path, scan_rows: int = 20) -> list[dict[str, object]]:
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    result = []
    for sheet in workbook.worksheets:
        candidates = []
        for row_number in range(1, min(sheet.max_row, scan_rows) + 1):
            values = [clean_spaces(cell.value) for cell in sheet[row_number]]
            nonempty = [value for value in values if value]
            if len(nonempty) >= 2 or any(normalize_header(value) == "ALAMAT" for value in nonempty):
                candidates.append({"row": row_number, "columns": nonempty})
        result.append(
            {
                "sheet": sheet.title,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "header_candidates": candidates[:5],
            }
        )
    workbook.close()
    return result
