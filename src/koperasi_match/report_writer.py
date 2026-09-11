from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import MatchResult

AUDIT_HEADERS = [
    "ALAMAT_SUMBER",
    "SOURCE_SHEET",
    "SOURCE_ROW",
    "DESA_PARSED",
    "KECAMATAN_PARSED",
    "KABUPATEN_PARSED",
    "METODE_MATCH",
    "MATCH_SCORE",
    "MATCH_MARGIN",
    "REFERENCE_ROW",
]
REVIEW_HEADERS = [
    "ALAMAT_SUMBER",
    "SOURCE_SHEET",
    "SOURCE_ROW",
    "DESA_PARSED",
    "KECAMATAN_PARSED",
    "KABUPATEN_PARSED",
    "STATUS",
    "BEST_SCORE",
    "SECOND_SCORE",
    "SCORE_MARGIN",
    "BEST_REFERENCE_ROW",
    "BEST_NIK",
    "BEST_DESA",
    "BEST_KECAMATAN",
    "BEST_KABUPATEN",
    "BEST_PROVINSI",
    "CATATAN",
    "KANDIDAT_2",
    "KANDIDAT_3",
]
NOT_FOUND_HEADERS = [
    "ALAMAT_SUMBER",
    "SOURCE_SHEET",
    "SOURCE_ROW",
    "DESA_PARSED",
    "KECAMATAN_PARSED",
    "KABUPATEN_PARSED",
    "STATUS",
    "ALASAN",
    "KANDIDAT_DIPERIKSA",
]


def _parsed_values(result: MatchResult) -> list[Any]:
    parsed = result.source.parsed
    return [
        result.source.address,
        result.source.sheet,
        result.source.row,
        parsed.desa if parsed else "",
        parsed.kecamatan if parsed else "",
        parsed.kabupaten if parsed else "",
    ]


def _candidate_label(result: MatchResult, index: int) -> str:
    if index >= len(result.top_candidates):
        return ""
    candidate = result.top_candidates[index]
    record = candidate.reference
    location = f"{record.desa_original} | {record.kecamatan_original} | {record.kabupaten_original}"
    return f"{location} | {candidate.score:.2f}"


def _style_sheet(sheet, status_column: int | None = None) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in range(1, sheet.max_column + 1):
        values = [
            str(sheet.cell(row, column).value or "")
            for row in range(1, min(sheet.max_row, 200) + 1)
        ]
        width = min(max(max((len(value) for value in values), default=8) + 2, 10), 45)
        sheet.column_dimensions[get_column_letter(column)].width = width
    if status_column:
        fills = {
            "MATCHED": "E2F0D9",
            "REVIEW": "FFF2CC",
            "NOT_FOUND": "FCE4D6",
            "PARSE_FAILED": "F4CCCC",
        }
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, status_column)
            for prefix, color in fills.items():
                if str(cell.value or "").startswith(prefix):
                    cell.fill = PatternFill("solid", fgColor=color)
                    break
    sheet.sheet_view.showGridLines = False


def write_report(
    output: Path,
    reference_headers: list[object],
    results: list[MatchResult],
    source: Path,
    reference: Path,
    sheets: list[str],
    threshold: float,
    minimum_margin: float,
    duplicate_reference_count: int,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    matched = workbook.create_sheet("MATCHED")
    review = workbook.create_sheet("PERLU_REVIEW")
    not_found = workbook.create_sheet("TIDAK_DITEMUKAN")
    summary = workbook.create_sheet("RINGKASAN", 0)

    matched.append([*reference_headers, *AUDIT_HEADERS])
    review.append(REVIEW_HEADERS)
    not_found.append(NOT_FOUND_HEADERS)

    for result in results:
        if result.is_matched and result.reference:
            matched.append(
                [
                    *result.reference.values,
                    *_parsed_values(result),
                    result.status.value,
                    result.score,
                    result.margin,
                    result.reference.row,
                ]
            )
        elif result.status.value.startswith("REVIEW"):
            best = result.top_candidates[0].reference if result.top_candidates else None
            review.append(
                [
                    *_parsed_values(result),
                    result.status.value,
                    result.score,
                    result.second_score,
                    result.margin,
                    best.row if best else "",
                    best.nik if best else "",
                    best.desa_original if best else "",
                    best.kecamatan_original if best else "",
                    best.kabupaten_original if best else "",
                    best.province_original if best else "",
                    result.note,
                    _candidate_label(result, 1),
                    _candidate_label(result, 2),
                ]
            )
        else:
            not_found.append(
                [
                    *_parsed_values(result),
                    result.status.value,
                    result.note,
                    result.candidates_checked,
                ]
            )

    counts = Counter(result.status.value for result in results)
    summary_rows = [
        ("Laporan pencocokan koperasi", ""),
        ("File sumber", str(source)),
        ("File referensi", str(reference)),
        ("Sheet sumber", ", ".join(sheets)),
        ("Waktu pemrosesan", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Threshold fuzzy", threshold),
        ("Minimum margin", minimum_margin),
        ("Jumlah baris sumber", len(results)),
        (
            "Jumlah matched",
            sum(value for key, value in counts.items() if key.startswith("MATCHED_")),
        ),
        (
            "Jumlah perlu review",
            sum(value for key, value in counts.items() if key.startswith("REVIEW_")),
        ),
        ("Jumlah tidak ditemukan", counts["NOT_FOUND"]),
        ("Jumlah parse failed", counts["PARSE_FAILED"]),
        ("Duplikasi key referensi", duplicate_reference_count),
    ]
    summary.append(["METRIK", "NILAI"])
    for row in summary_rows:
        summary.append(row)
    summary.append([])
    summary.append(["STATUS", "JUMLAH"])
    for status, count in sorted(counts.items()):
        summary.append([status, count])

    for sheet in (matched, review, not_found):
        status_header = "METODE_MATCH" if sheet.title == "MATCHED" else "STATUS"
        status_column = next(
            (cell.column for cell in sheet[1] if cell.value == status_header), None
        )
        _style_sheet(sheet, status_column)
    _style_sheet(summary)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 80

    for sheet in (matched, review):
        for row in range(2, sheet.max_row + 1):
            for header in ("NIK", "pic phone"):
                for cell in sheet[1]:
                    if str(cell.value or "").strip().lower() == header.lower():
                        sheet.cell(row, cell.column).number_format = "@"

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
